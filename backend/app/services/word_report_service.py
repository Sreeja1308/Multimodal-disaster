from __future__ import annotations

from io import BytesIO
from typing import Any

from docx import Document
from docx.shared import Inches, Pt


class WordReportService:
    def build_docx(self, report: dict[str, Any]) -> BytesIO:
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

        title = document.add_heading(report.get("report_type", "FloodWatch Disaster Report"), 0)
        title.runs[0].font.size = Pt(22)
        document.add_paragraph(f"Generated: {report.get('generated_at', 'Unavailable')}")

        location = report.get("location", {})
        document.add_heading("Selected location", level=1)
        document.add_paragraph(f"Latitude: {location.get('latitude', 'Unavailable')}\nLongitude: {location.get('longitude', 'Unavailable')}")

        risk = report.get("risk_assessment", {})
        document.add_heading("Risk assessment", level=1)
        document.add_paragraph(
            f"Risk level: {risk.get('risk_level', 'Unavailable')}\n"
            f"Prediction status: {risk.get('prediction_status', 'Unavailable')}\n"
            f"Probability: {risk.get('probability', 'Unavailable')}\n"
            f"Model version: {risk.get('model_version', 'Unavailable')}"
        )
        for limitation in risk.get("limitations", []):
            document.add_paragraph(limitation, style="List Bullet")

        environment = report.get("observed_environment", {})
        document.add_heading("Observed environmental data", level=1)
        document.add_paragraph(
            f"Provider: {environment.get('provider', 'Unavailable')}\n"
            f"Observed at: {environment.get('observed_at', 'Unavailable')}\n"
            f"Retrieved at: {environment.get('retrieved_at', 'Unavailable')}"
        )
        for name, value in (environment.get("values") or {}).items():
            document.add_paragraph(f"{name.replace('_', ' ').title()}: {value if value is not None else 'Unavailable'}", style="List Bullet")
        for name, reason in (environment.get("unavailable") or {}).items():
            document.add_paragraph(f"{name.replace('_', ' ').title()}: Unavailable ({reason})", style="List Bullet")

        document.add_heading("Forecast and limitations", level=1)
        forecast = report.get("forecast", {})
        document.add_paragraph(f"Status: {forecast.get('status', 'Unavailable')}\n{forecast.get('note', '')}")
        for day in forecast.get("days", []):
            document.add_paragraph(
                f"{day.get('date', 'Unavailable')}: {day.get('temperature_min_c', 'Unavailable')}°C to "
                f"{day.get('temperature_max_c', 'Unavailable')}°C; precipitation "
                f"{day.get('precipitation_mm', 'Unavailable')} mm; probability "
                f"{day.get('precipitation_probability_pct', 'Unavailable')}%",
                style="List Bullet",
            )

        document.add_heading("Current incidents", level=1)
        incidents = report.get("incidents") or []
        if incidents:
            table = document.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            for cell, text in zip(table.rows[0].cells, ["Name", "Location", "Severity", "Status", "Source"]):
                cell.text = text
            for incident in incidents:
                cells = table.add_row().cells
                for cell, key in zip(cells, ["name", "location", "severity", "status", "source"]):
                    cell.text = str(incident.get(key, "Unavailable"))
        else:
            document.add_paragraph("No incidents are currently recorded.")

        document.add_heading("Recommended precautions", level=1)
        for action in report.get("recommended_precautions", []):
            document.add_paragraph(action, style="List Bullet")

        document.add_heading("Data sources", level=1)
        document.add_paragraph(str(report.get("data_sources", {})))

        output = BytesIO()
        document.save(output)
        output.seek(0)
        return output
